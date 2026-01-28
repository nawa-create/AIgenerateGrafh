"""レポート出力モジュール（PDF/PPTX）"""

import io
from datetime import datetime


def generate_pdf_report(
    charts: list[dict],
    data_profile: dict | None = None,
    title: str = "分析レポート",
    subtitle: str = "AI Graph Generator により自動生成",
) -> bytes:
    """
    PDF形式の分析レポートを生成する

    Args:
        charts: グラフ情報のリスト [{"title/proposal", "figure", "code", "insights"}]
        data_profile: データプロファイル
        title: レポートタイトル
        subtitle: サブタイトル

    Returns:
        PDFのバイナリデータ
    """
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, Table, TableStyle, PageBreak
    from reportlab.lib import colors
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    import tempfile
    import os

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=20*mm, bottomMargin=20*mm)
    styles = getSampleStyleSheet()

    # Try to register Japanese font (fall back to Helvetica if not available)
    japanese_font = "Helvetica"
    for font_path in [
        "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/noto-cjk/NotoSansCJKjp-Regular.otf",
        "/System/Library/Fonts/ヒラギノ角ゴシック W3.ttc",
    ]:
        if os.path.exists(font_path):
            try:
                pdfmetrics.registerFont(TTFont("JapaneseFont", font_path))
                japanese_font = "JapaneseFont"
                break
            except Exception:
                continue

    title_style = ParagraphStyle(
        "CustomTitle", parent=styles["Title"],
        fontName=japanese_font, fontSize=24, spaceAfter=10
    )
    subtitle_style = ParagraphStyle(
        "CustomSubtitle", parent=styles["Normal"],
        fontName=japanese_font, fontSize=12, textColor=colors.grey, spaceAfter=30
    )
    heading_style = ParagraphStyle(
        "CustomHeading", parent=styles["Heading2"],
        fontName=japanese_font, fontSize=16, spaceBefore=20, spaceAfter=10
    )
    body_style = ParagraphStyle(
        "CustomBody", parent=styles["Normal"],
        fontName=japanese_font, fontSize=10, spaceAfter=6
    )
    insight_style = ParagraphStyle(
        "InsightStyle", parent=styles["Normal"],
        fontName=japanese_font, fontSize=10, spaceAfter=4,
        leftIndent=10, bulletIndent=0
    )

    elements = []

    # Title page
    elements.append(Spacer(1, 50*mm))
    elements.append(Paragraph(title, title_style))
    elements.append(Paragraph(subtitle, subtitle_style))
    elements.append(Paragraph(
        f"生成日時: {datetime.now().strftime('%Y年%m月%d日 %H:%M')}",
        body_style
    ))

    # Data profile summary
    if data_profile and "files" in data_profile:
        elements.append(Spacer(1, 20*mm))
        elements.append(Paragraph("データ概要", heading_style))
        for f in data_profile["files"]:
            name = f.get("name", "不明")
            rows = f.get("rows", "?")
            cols = len(f.get("columns", []))
            elements.append(Paragraph(
                f"・{name}（{rows}行 × {cols}列）", body_style
            ))

    elements.append(PageBreak())

    # Charts section
    for i, chart in enumerate(charts):
        chart_title = ""
        if "proposal" in chart and isinstance(chart["proposal"], dict):
            chart_title = chart["proposal"].get("title", f"グラフ {i+1}")
        elif "title" in chart:
            chart_title = chart["title"]
        else:
            chart_title = f"グラフ {i+1}"

        elements.append(Paragraph(f"{i+1}. {chart_title}", heading_style))

        # Export chart as image
        fig = chart.get("figure")
        if fig is not None:
            try:
                img_bytes = fig.to_image(format="png", width=700, height=450, scale=2)
                img_buffer = io.BytesIO(img_bytes)
                img = Image(img_buffer, width=170*mm, height=110*mm)
                elements.append(img)
                elements.append(Spacer(1, 5*mm))
            except Exception:
                elements.append(Paragraph("（グラフ画像の出力に失敗しました）", body_style))

        # Insights
        insights_data = chart.get("insights") or chart.get("insight")
        if insights_data and isinstance(insights_data, dict):
            insight_list = insights_data.get("insights", [])
            if insight_list:
                elements.append(Paragraph("分析インサイト:", body_style))
                for ins in insight_list:
                    severity = ins.get("severity", "info")
                    marker = "⚠️" if severity == "warning" else "💡"
                    msg = ins.get("message", "")
                    elements.append(Paragraph(f"{marker} {msg}", insight_style))

            recs = insights_data.get("recommendations", [])
            if recs:
                elements.append(Spacer(1, 3*mm))
                elements.append(Paragraph("推奨アクション:", body_style))
                for rec in recs:
                    elements.append(Paragraph(f"✅ {rec}", insight_style))

        if i < len(charts) - 1:
            elements.append(PageBreak())

    doc.build(elements)
    return buffer.getvalue()


def generate_pptx_report(
    charts: list[dict],
    data_profile: dict | None = None,
    title: str = "分析レポート",
    subtitle: str = "AI Graph Generator により自動生成",
) -> bytes:
    """
    PPTX形式の分析レポートを生成する

    Args:
        charts: グラフ情報のリスト
        data_profile: データプロファイル
        title: レポートタイトル
        subtitle: サブタイトル

    Returns:
        PPTXのバイナリデータ
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Title slide
    slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(slide_layout)

    # Title text box
    left = Inches(1)
    top = Inches(2)
    width = Inches(11)
    height = Inches(1.5)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(40)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

    # Subtitle
    p2 = tf.add_paragraph()
    p2.text = subtitle
    p2.font.size = Pt(18)
    p2.font.color.rgb = RGBColor(128, 128, 128)
    p2.alignment = PP_ALIGN.CENTER

    # Date
    p3 = tf.add_paragraph()
    p3.text = f"生成日時: {datetime.now().strftime('%Y年%m月%d日 %H:%M')}"
    p3.font.size = Pt(14)
    p3.font.color.rgb = RGBColor(160, 160, 160)
    p3.alignment = PP_ALIGN.CENTER

    # Chart slides
    for i, chart in enumerate(charts):
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        chart_title = ""
        if "proposal" in chart and isinstance(chart["proposal"], dict):
            chart_title = chart["proposal"].get("title", f"グラフ {i+1}")
        elif "title" in chart:
            chart_title = chart["title"]
        else:
            chart_title = f"グラフ {i+1}"

        # Title
        txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.8))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = chart_title
        p.font.size = Pt(28)
        p.font.bold = True

        # Chart image
        fig = chart.get("figure")
        if fig is not None:
            try:
                img_bytes = fig.to_image(format="png", width=1200, height=700, scale=2)
                img_stream = io.BytesIO(img_bytes)
                slide.shapes.add_picture(
                    img_stream, Inches(0.5), Inches(1.2), Inches(8.5), Inches(5)
                )
            except Exception:
                txBox2 = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(6), Inches(1))
                txBox2.text_frame.paragraphs[0].text = "（グラフ画像の出力に失敗しました）"

        # Insights on right side
        insights_data = chart.get("insights") or chart.get("insight")
        if insights_data and isinstance(insights_data, dict):
            insight_list = insights_data.get("insights", [])
            recs = insights_data.get("recommendations", [])

            if insight_list or recs:
                txBox3 = slide.shapes.add_textbox(
                    Inches(9.2), Inches(1.2), Inches(3.8), Inches(5.5)
                )
                tf3 = txBox3.text_frame
                tf3.word_wrap = True

                if insight_list:
                    p = tf3.paragraphs[0]
                    p.text = "💡 インサイト"
                    p.font.size = Pt(14)
                    p.font.bold = True

                    for ins in insight_list:
                        p_ins = tf3.add_paragraph()
                        severity = ins.get("severity", "info")
                        marker = "⚠️" if severity == "warning" else "•"
                        p_ins.text = f"{marker} {ins.get('message', '')}"
                        p_ins.font.size = Pt(11)
                        p_ins.space_after = Pt(4)

                if recs:
                    p_rec_title = tf3.add_paragraph()
                    p_rec_title.text = ""
                    p_rec_title = tf3.add_paragraph()
                    p_rec_title.text = "✅ 推奨アクション"
                    p_rec_title.font.size = Pt(14)
                    p_rec_title.font.bold = True

                    for rec in recs:
                        p_rec = tf3.add_paragraph()
                        p_rec.text = f"• {rec}"
                        p_rec.font.size = Pt(11)
                        p_rec.space_after = Pt(4)

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()
